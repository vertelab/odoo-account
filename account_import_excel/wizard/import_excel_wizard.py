import base64
from openpyxl import load_workbook
from io import BytesIO

from odoo import models, fields, api, _
from odoo.exceptions import UserError, AccessError, ValidationError
import logging

_logger = logging.getLogger(__name__)

class ImportExcelWizard(models.TransientModel):
    _name = 'import.excel.wizard'
    _description = 'Simpel excel import wizard'

    excel_file = fields.Binary(required=True)

    def load_excel(self):

        excel_file = BytesIO(base64.b64decode(self.excel_file))
        try:
            wb = load_workbook(excel_file)
        except Exception as e:
            raise UserError(f"The file given could not be read as an Excel file!\n\n{e}")
        
        header_fields = {"invoicenumber": "ref","date": "invoice_date_due","journal": "journal_id"}
        header_field_rows = {
            "contact": "partner_id", 
            "product": "product_id", 
            "account":"account_id", 
            "unit": "product_uom_id", 
            "amount": "quantity", 
            "price": "price_unit", 
            "tax": "tax_ids"} #"analytic_accounts": "analytic_distribution",

        for sheet in wb:
            first_row = next(sheet.iter_rows(values_only=True,min_row=1,max_row=1))
            move_id = self.create_account_move(header_fields,sheet,first_row)
            self.create_account_move_line(header_field_rows,sheet,first_row,move_id)
           
    def create_account_move(self,header_fields,sheet,first_row):
        header_field_index,filterd_header_fields = self.index_word_filter(first_row,header_fields)
        header_field_values = self.get_header_field_values(sheet,header_field_index,filterd_header_fields)
        data = self._update_move_values(header_field_values)
        if journal_id := data.get("journal_id"):
              self._set_types(data,journal_id)
        _logger.error(f"{data=}")
        move_id = self.env["account.move"].create(data)
        return move_id
    
    def create_account_move_line(self,header_field_rows,sheet,first_row,move_id):
        word_index,filterd_row_fields = self.index_word_filter(first_row,header_field_rows)
        header_field_row_values = self.get_header_field_row_values(sheet,word_index,filterd_row_fields)
        move_line_ids = []
        for data in header_field_row_values:
            data.update({"move_id":move_id.id})
            data = self._update_move_values(data)
            _logger.error(f"{data=}")
            move_line_ids.append(self.env["account.move.line"].create(data))
        return move_line_ids

    def _update_move_values(self,data:dict):
        for key,value in data.items():
            new_value = value
            if key == "partner_id" and isinstance(value,str):
                new_value = self._is_string("res.partner",value)
            if key == "product_id" and isinstance(value,str):
                new_value = self._is_string("product.product",value)
            if key == "analytic_accounts_id" and isinstance(value,str):
                new_value = self._is_string("account.analytic.account",value)
            if key == "product_uom_id" and isinstance(value,str):
                new_value = self._is_string("uom.uom",value)
            if key == "account_id":
                new_value = self.is_account_code(value)
            if key == "journal_id" and isinstance(value,str):
                new_value = self._is_string("account.journal",value)
            if key == "tax_ids" and "%" in value:
                new_value = [(4,self._is_string("account.tax",value))]
            data.update({key:new_value})
        return data
    
    def _set_types(self,data,journal_id):
        journal_id = self.env["account.journal"].browse(journal_id)
        if journal_id.type == "sale":
            data.update({
                "move_type": "out_invoice",
                }) 
        elif journal_id.type == "purchase":
            data.update({
                "move_type": "in_invoice",
                })
        return data
    
    def is_account_code(self,value):
        new_value = value
        model_id = False
        if isinstance(value,str): 
            value = value.strip()
            ext_id = self.env.ref(value,raise_if_not_found=False)
            if ext_id:
                new_value = ext_id.id
        else:
            model_id = self.env["account.account"].search([
                "|",
                ("name","ilike",value),
                ("code","=",value)],
                limit=1)
        if model_id:
            new_value = model_id.id
        return new_value

    def _is_string(self,model,value):
        value = value.strip()
        new_value = value
        ext_id = self.env.ref(value,raise_if_not_found=False)
        model_id = self.env[model].search([
            "|",
            ("name","ilike",value),
            ("display_name","=",value)],
            limit=1)
        if ext_id:
            new_value = ext_id.id
        elif model_id:
            new_value = model_id.id
        return new_value

    def get_header_field_row_values(self,sheet,word_index,filterd_row_fields):
        datas = []
        for row in sheet.iter_rows(values_only=True,min_row=2,max_row=100,max_col=100):
            if not all(r is None for r in row):
                row_values = [row[i] for i in word_index]
                datas.append(dict(zip(filterd_row_fields,row_values)))
        return datas

    def get_header_field_values(self,sheet,header_field_index,filterd_header_fields):
        while True:
            row = next(sheet.iter_rows(values_only=True,min_row=2,max_row=100,max_col=100))
            if not all(r is None for r in row):
                row_values = [row[i] for i in header_field_index]
                return dict(zip(filterd_header_fields,row_values))

    def index_word_filter(self,first_row,words):
        word_index = []
        filterd_words_fields = []
        for word in words.keys():
            if word in first_row:
                row_index = first_row.index(word)
                word_index.append(row_index)
                filterd_words_fields.append(words[word])
        return word_index, filterd_words_fields

